#!/usr/bin/env python3
import json
import os
import sys
import openpyxl

def main():
    excel_path = os.path.join(os.path.dirname(__file__), '..', 'SAOS_Health_Rules_Tracker_v2_9574 1.xlsx')
    output_path = os.path.join(os.path.dirname(__file__), '..', 'src', 'lib', 'rules', 'master-catalog.json')
    
    if not os.path.exists(excel_path):
        print(f"Error: Excel file not found at {excel_path}", file=sys.stderr)
        sys.exit(1)
        
    print(f"Loading workbook from {excel_path}...")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    
    sheet_mapping = {
        'CMDB': 'CMDB',
        'ITSM': 'ITSM',
        'ITOM': 'ITOM',
        'platform': 'Platform',
        'data quality': 'Data Quality'
    }
    
    all_rules = []
    rules_by_domain = {}
    
    for sheet_name, domain_name in sheet_mapping.items():
        if sheet_name not in wb.sheetnames:
            print(f"Warning: Sheet {sheet_name} not found in workbook", file=sys.stderr)
            continue
            
        sheet = wb[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows or len(rows) < 2:
            continue
            
        is_dq = (sheet_name == 'data quality')
        domain_rules = []
        
        for row_idx, r in enumerate(rows[1:], start=2):
            if not r or r[0] is None:
                continue
            rule_id = str(r[0]).strip()
            if not rule_id or rule_id.startswith('CSDM') or rule_id.lower() == 'rule id':
                continue
                
            group = str(r[2] if is_dq else r[1] or '').strip()
            title = str(r[3] if is_dq else r[2] or '').strip()
            base_severity = str(r[4] if is_dq else r[3] or 'Moderate').strip()
            what_it_means = str(r[5] if is_dq else r[4] or '').strip()
            why_it_matters = str(r[6] if is_dq else r[5] or '').strip()
            source_tables = str(r[7] if is_dq else r[6] or '').strip()
            detection_logic = str(r[8] if is_dq else r[7] or '').strip()
            threshold = str(r[9] if is_dq else r[8] or '').strip()
            confidence_basis = str(r[10] if is_dq else r[9] or '').strip()
            evidence_to_show = str(r[11] if is_dq else r[10] or '').strip()
            false_positive_guard = str(r[12] if is_dq else r[11] or '').strip()
            remediation_lane = str(r[13] if is_dq else r[12] or 'Lane 2').strip()
            cross_domain_link = str(r[14] if is_dq else r[13] or '').strip()
            
            # Normalize severity to standard 5 tiers
            norm_severity = "Moderate"
            sev_lower = base_severity.lower()
            if "systemic" in sev_lower:
                norm_severity = "Systemic"
            elif "critical" in sev_lower:
                norm_severity = "Critical"
            elif "high" in sev_lower:
                norm_severity = "High"
            elif "moderate" in sev_lower or "medium" in sev_lower:
                norm_severity = "Moderate"
            elif "low" in sev_lower:
                norm_severity = "Low"
                
            # Normalize lane
            norm_lane = 2
            lane_lower = remediation_lane.lower()
            if "lane 1" in lane_lower:
                norm_lane = 1
            elif "lane 3" in lane_lower:
                norm_lane = 3
            else:
                norm_lane = 2
                
            rule_item = {
                "id": rule_id,
                "domain": domain_name,
                "group": group,
                "title": title,
                "baseSeverity": norm_severity,
                "rawSeverity": base_severity,
                "whatItMeans": what_it_means,
                "whyItMatters": why_it_matters,
                "sourceTables": source_tables,
                "detectionLogic": detection_logic,
                "threshold": threshold,
                "confidenceBasis": confidence_basis,
                "evidenceToShow": evidence_to_show,
                "falsePositiveGuard": false_positive_guard,
                "remediationLane": norm_lane,
                "rawRemediationLane": remediation_lane,
                "crossDomainLink": cross_domain_link,
            }
            domain_rules.append(rule_item)
            all_rules.append(rule_item)
            
        rules_by_domain[domain_name] = len(domain_rules)
        print(f"Extracted {len(domain_rules)} rules from sheet '{sheet_name}' ({domain_name})")
        
    catalog_data = {
        "metadata": {
            "version": "2.0",
            "sourceFile": "SAOS_Health_Rules_Tracker_v2_9574 1.xlsx",
            "totalRules": len(all_rules),
            "domainCounts": rules_by_domain,
            "generatedAt": "2026-09-20T18:05:00Z"
        },
        "rules": all_rules
    }
    
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(catalog_data, f, indent=2, ensure_ascii=False)
        
    print(f"Successfully wrote {len(all_rules)} rules to {output_path}")

if __name__ == '__main__':
    main()
